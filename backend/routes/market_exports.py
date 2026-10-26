"""Bounded, read-only exports of public market datasets."""
from __future__ import annotations

import csv
import io
import logging
import os
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Callable

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook

from backend import db_path
from backend.sqlite_utils import read_connection
from backend.routes.downloads import rate_limit_download

logger = logging.getLogger(__name__)
market_exports_bp = Blueprint("market_exports", __name__, url_prefix="/api/exports")
MAX_ROWS = 100_000
ROOT = Path(__file__).resolve().parents[2]


@market_exports_bp.after_request
def no_export_cache(response):
    response.headers['Cache-Control'] = 'no-store'
    return response


def local_db(filename: str, env: str = "") -> Callable[[], str]:
    def resolve() -> str:
        if env and os.getenv(env):
            return os.environ[env]
        relative = Path('data') / ('funds.sqlite' if filename == 'funds.sqlite' else f'sqlite/{filename}')
        local = ROOT / relative
        return str(local if local.exists() else Path('/var/www/valuation') / relative)
    return resolve


@dataclass(frozen=True)
class Dataset:
    vi: str
    en: str
    path: Callable[[], str]
    table: str
    columns: str
    date_column: str = ""
    symbol_column: str = ""
    require_symbols: bool = False
    source: str = "Vietcap"
    note_vi: str = "Giữ nguyên giá trị và đơn vị của nguồn; ô trống là dữ liệu chưa có."
    note_en: str = "Original source values and units; blank cells indicate missing data."

    @property
    def fields(self) -> list[str]:
        return self.columns.split(',')


DATASETS = {
    'index-valuations': Dataset(
        'P/E, P/B và giá VNINDEX',
        'VNINDEX valuations and prices',
        db_path.resolve_vci_valuation_db_path,
        'valuation_history',
        'date,pe,pb,vnindex,open,high,low,close,volume,ema50',
        'date',
    ),
    'index-prices': Dataset(
        'Lịch sử chỉ số thị trường',
        'Market index history',
        db_path.resolve_index_history_db_path,
        'market_index_history',
        'symbol,tradingDate,openIndex,highestIndex,lowestIndex,closeIndex,totalVolume,totalValue,indexChange,percentIndexChange',
        'tradingDate',
        'symbol',
    ),
    'stock-prices': Dataset(
        'Lịch sử giá cổ phiếu',
        'Stock price history',
        db_path.resolve_price_history_db_path,
        'stock_price_history',
        'symbol,time,open,high,low,close,volume',
        'time',
        'symbol',
        True,
        note_vi='OHLC gốc từ kho Vietcap; có thể khác giá đã điều chỉnh trên biểu đồ. Giữ nguyên đơn vị nguồn; ô trống là dữ liệu chưa có.',
        note_en='Raw Vietcap OHLC may differ from adjusted chart prices. Original source units; blank cells indicate missing data.',
    ),
    'daily-ratios': Dataset(
        'P/E, P/B cổ phiếu theo ngày',
        'Daily stock PE and PB',
        db_path.resolve_vci_ratio_daily_db_path,
        'ratio_daily_history',
        'ticker,trading_date,pe,pb',
        'trading_date',
        'ticker',
    ),
    'companies': Dataset(
        'Danh mục doanh nghiệp và ngành',
        'Company and sector directory',
        db_path.resolve_vci_company_db_path,
        'companies',
        'ticker,organ_name,en_organ_name,floor,icb_code1,icb_name1,icb_code2,icb_name2,icb_code3,icb_name3,icb_code4,icb_name4,fetched_at',
        symbol_column='ticker',
    ),
    'screener': Dataset(
        'Bộ lọc cổ phiếu — bản mới nhất',
        'Latest stock screener snapshot',
        db_path.resolve_vci_screening_db_path,
        'screening_data',
        'ticker,exchange,marketPrice,marketCap,accumulatedVolume,accumulatedValue,dailyPriceChangePercent,ttmPe,ttmPb,ttmRoe,netMargin,grossMargin,revenueGrowthYoy,npatmiGrowthYoyQm1,viSector,enSector,fetched_at',
        symbol_column='ticker',
    ),
    'foreign-flow': Dataset(
        'Giao dịch khối ngoại theo phút',
        'Foreign trading by minute',
        local_db('vci_foreign.sqlite', 'VCI_FOREIGN_DB_PATH'),
        'foreign_volume_minute',
        'trading_date,minute,buy_volume,sell_volume,buy_value,sell_value',
        'trading_date',
    ),
    'macro-prices': Dataset(
        'Vĩ mô Việt Nam và giá quốc tế',
        'Vietnam macro and world prices',
        local_db('macro_history.sqlite'),
        'macro_prices',
        'symbol,date,close',
        'date',
        'symbol',
        source='Yahoo Finance / TradingView',
        note_vi='Mã ECONOMICS:... là chỉ tiêu kinh tế từ TradingView; các mã khác là giá/FX từ Yahoo Finance. Đơn vị tùy từng mã.',
        note_en='ECONOMICS:... symbols are TradingView economic indicators; other symbols are Yahoo Finance prices/FX. Units vary by symbol.',
    ),
    'news-events': Dataset(
        'Tin tức và sự kiện doanh nghiệp',
        'Company news and events',
        db_path.resolve_vci_news_events_db_path,
        'items',
        'symbol,tab,public_date,title',
        'public_date',
        'symbol',
    ),
    'fund-nav': Dataset(
        'Lịch sử NAV quỹ',
        'Fund NAV history',
        local_db('funds.sqlite', 'FUNDS_DB_PATH'),
        'fund_nav',
        'symbol,date,nav',
        'date',
        'symbol',
        source='Nguồn quỹ / fund providers',
    ),
    'market-breadth': Dataset(
        'Độ rộng thị trường trên EMA50',
        'Market breadth above EMA50',
        db_path.resolve_vci_valuation_db_path,
        'ema_breadth_history',
        'trading_date,above_count,total_count,above_percent',
        'trading_date',
    ),
}


def _available(dataset: Dataset) -> bool:
    try:
        with read_connection(dataset.path(), timeout=1) as conn:
            # A dataset is offered only when the actual schema and rows exist.
            return conn.execute(f'SELECT {dataset.columns} FROM {dataset.table} LIMIT 1').fetchone() is not None
    except (sqlite3.Error, OSError):
        return False


@market_exports_bp.get('/catalog')
def catalog():
    return jsonify({'max_rows': MAX_ROWS, 'datasets': [
        {'id': key, 'name': {'vi': ds.vi, 'en': ds.en}, 'available': _available(ds),
         'columns': ds.fields, 'dated': bool(ds.date_column),
         'symbols': bool(ds.symbol_column), 'require_symbols': ds.require_symbols,
         'source': ds.source, 'note': {'vi': ds.note_vi, 'en': ds.note_en}}
        for key, ds in DATASETS.items()
    ]})


def export_query(dataset: Dataset, args) -> tuple[str, list]:
    """Only identifiers from the fixed dataset registry can enter SQL."""
    where, params = [], []
    dates = []
    for key, operator in [('from', '>='), ('to', '<')]:
        raw = args.get(key, '')
        if not raw:
            dates.append(None)
            continue
        if not dataset.date_column:
            raise ValueError('This dataset is a latest snapshot; date filters are unsupported.')
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', raw):
            raise ValueError('Dates must use YYYY-MM-DD.')
        parsed = date.fromisoformat(raw)
        dates.append(parsed)
        # ISO timestamps also match the complete inclusive final day.
        if key == 'to':
            raw = (parsed + timedelta(days=1)).isoformat()
        where.append(f'{dataset.date_column} {operator} ?')
        params.append(raw)
    if all(dates) and dates[0] > dates[1]:
        raise ValueError('The start date must not exceed the end date.')
    symbols = [s.strip().upper() for s in args.get('symbols', '').split(',') if s.strip()]
    if dataset.require_symbols and not symbols:
        raise ValueError('Select at least one symbol (up to 50).')
    if symbols:
        if not dataset.symbol_column:
            raise ValueError('Symbol filters are unsupported for this dataset.')
        if len(symbols) > 50 or any(not re.fullmatch(r'[A-Z0-9^][A-Z0-9.^=:_-]{0,39}', s) for s in symbols):
            raise ValueError('Invalid symbols; select up to 50 symbols.')
        # NOCASE permits provider-specific index spelling (e.g. HNXIndex).
        where.append(f'{dataset.symbol_column} COLLATE NOCASE IN ({",".join("?" for _ in symbols)})')
        params.extend(symbols)
    order = ','.join(c for c in (dataset.symbol_column, dataset.date_column) if c)
    sql = f'SELECT {dataset.columns} FROM {dataset.table}'
    if where:
        sql += ' WHERE ' + ' AND '.join(where)
    if order:
        sql += ' ORDER BY ' + order
    return sql + ' LIMIT ?', params + [MAX_ROWS + 1]


def safe_cell(value):
    # Spreadsheet formula injection protection applies to text, not negative numbers.
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
        return "'" + value
    return value


@market_exports_bp.get('/market')
@rate_limit_download
def export_market():
    dataset_id = request.args.get('dataset', '')
    dataset = DATASETS.get(dataset_id)
    if dataset is None:
        return jsonify({'error': 'Unknown dataset.'}), 400
    fmt = request.args.get('format', 'csv')
    if fmt not in {'csv', 'xlsx'}:
        return jsonify({'error': 'Format must be csv or xlsx.'}), 400
    try:
        sql, params = export_query(dataset, request.args)
    except (ValueError, OverflowError):
        return jsonify({'error': 'Invalid filters. Use YYYY-MM-DD, an ordered date range and up to 50 valid symbols; stock prices require symbols.'}), 400
    try:
        with read_connection(dataset.path()) as conn:
            rows = conn.execute(sql, params).fetchall()
    except (sqlite3.Error, OSError):
        logger.exception('Market export unavailable: %s', dataset_id)
        return jsonify({'error': 'This dataset is temporarily unavailable.'}), 503
    if not rows:
        return jsonify({'error': 'No data matches these filters.'}), 404
    if len(rows) > MAX_ROWS:
        return jsonify({'error': f'The selection exceeds {MAX_ROWS:,} rows. Narrow the date range or symbols; no data was truncated.'}), 413
    output = io.BytesIO()
    if fmt == 'csv':
        text = io.StringIO(newline='')
        writer = csv.writer(text)
        writer.writerow(dataset.fields)
        writer.writerows(tuple(safe_cell(v) for v in row) for row in rows)
        output.write(text.getvalue().encode('utf-8-sig'))
        mimetype = 'text/csv'
    else:
        book = Workbook(write_only=True)
        sheet = book.create_sheet('Data')
        sheet.append(dataset.fields)
        for row in rows:
            sheet.append([safe_cell(v) for v in row])
        notes = book.create_sheet('About')
        for row in [('Dataset', dataset.en), ('Source', dataset.source),
                    ('Exported at (UTC)', datetime.now(timezone.utc).isoformat()),
                    ('Rows', len(rows)), ('Filters', request.query_string.decode('utf-8', errors='replace')),
                    ('Note', dataset.note_en)]:
            notes.append([safe_cell(v) for v in row])
        book.save(output)
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    output.seek(0)
    response = send_file(output, mimetype=mimetype, as_attachment=True,
                         download_name=f'{dataset_id}-{date.today().isoformat()}.{fmt}')
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Export-Rows'] = str(len(rows))
    return response
