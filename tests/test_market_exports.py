"""Exports preserve values, reject oversized queries and never expose arbitrary SQL."""
import csv
import io
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from backend.routes import market_exports as exports
from backend.routes.downloads import download_tracker


class MarketExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = str(Path(self.tmp.name) / 'public.sqlite')
        with sqlite3.connect(self.path) as conn:
            conn.execute('CREATE TABLE prices (symbol TEXT, date TEXT, close REAL, title TEXT)')
            conn.executemany('INSERT INTO prices VALUES (?,?,?,?)', [
                ('FPT', '2026-10-22', 99.125678901, '=HYPERLINK("bad")'),
                ('FPT', '2026-10-23T15:00:00', -3.25, 'Tiếng Việt'),
                ('VCB', '2026-10-24', 40, 'Normal'),
            ])
        self.dataset = exports.Dataset('Giá', 'Prices', lambda: self.path, 'prices', 'symbol,date,close,title', 'date', 'symbol')
        self.patcher = patch.dict(exports.DATASETS, {'sample': self.dataset}, clear=True)
        self.patcher.start()
        download_tracker.clear()
        app = Flask(__name__)
        app.register_blueprint(exports.market_exports_bp)
        self.client = app.test_client()

    def tearDown(self):
        self.patcher.stop()
        self.tmp.cleanup()
        download_tracker.clear()

    def get(self, **params):
        return self.client.get('/api/exports/market', query_string={'dataset': 'sample', **params})

    def test_csv_inclusive_dates_preserves_precision_and_safe_text(self):
        response = self.get(symbols='fpt', **{'from': '2026-10-22', 'to': '2026-10-23'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['X-Export-Rows'], '2')
        rows = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'))))
        self.assertEqual(rows[0], self.dataset.fields)
        self.assertEqual(float(rows[1][2]), 99.125678901)
        self.assertTrue(rows[1][3].startswith("'="))
        self.assertEqual(rows[2][2:], ['-3.25', 'Tiếng Việt'])
        self.assertEqual(response.headers['Cache-Control'], 'no-store')

    def test_xlsx_preserves_numeric_and_unicode_values_and_documents_source(self):
        response = self.get(format='xlsx', symbols='FPT')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ['Data', 'About'])
        self.assertEqual(workbook['Data']['C2'].value, 99.125678901)
        self.assertNotEqual(workbook['Data']['D2'].data_type, 'f')
        self.assertEqual(workbook['Data']['C3'].value, -3.25)
        self.assertEqual(workbook['Data']['D3'].value, 'Tiếng Việt')
        self.assertIn(('Rows', 2), list(workbook['About'].values))
        workbook.close()

    def test_bad_filters_and_unregistered_dataset_are_rejected(self):
        cases = [{'dataset': '../../secrets'}, {'symbols': "FPT'); DROP TABLE prices;--"},
                 {'from': '2026-02-30'}, {'from': '2026-10-24', 'to': '2026-10-22'},
                 {'format': 'sqlite'}, {'symbols': ','.join(['FPT'] * 51)}]
        for params in cases:
            with self.subTest(params=params):
                self.assertEqual(self.get(**params).status_code, 400)
        with sqlite3.connect(self.path) as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM prices').fetchone()[0], 3)

    def test_empty_and_oversized_results_do_not_return_partial_files(self):
        self.assertEqual(self.get(symbols='NONE').status_code, 404)
        with patch.object(exports, 'MAX_ROWS', 1):
            response = self.get()
        self.assertEqual(response.status_code, 413)
        self.assertNotIn('Content-Disposition', response.headers)

    def test_missing_database_is_not_created_and_catalog_hides_it(self):
        missing = str(Path(self.tmp.name) / 'missing.sqlite')
        with patch.dict(exports.DATASETS, {'sample': replace(self.dataset, path=lambda: missing)}):
            self.assertFalse(self.client.get('/api/exports/catalog').json['datasets'][0]['available'])
            self.assertEqual(self.get().status_code, 503)
        self.assertFalse(Path(missing).exists())

    def test_stock_history_requires_symbols_and_snapshot_rejects_dates(self):
        with patch.dict(exports.DATASETS, {'sample': replace(self.dataset, require_symbols=True)}):
            self.assertEqual(self.get().status_code, 400)
            self.assertEqual(self.get(symbols='FPT').status_code, 200)
        with patch.dict(exports.DATASETS, {'sample': replace(self.dataset, date_column='')}):
            self.assertEqual(self.get(**{'from': '2026-10-01'}).status_code, 400)

    def test_catalog_exposes_only_public_metadata(self):
        item = self.client.get('/api/exports/catalog').json['datasets'][0]
        self.assertTrue(item['available'])
        self.assertEqual(item['columns'], self.dataset.fields)
        self.assertNotIn(self.path, str(item))
